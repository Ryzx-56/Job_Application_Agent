# utils/docx_generator.py
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from loguru import logger

OUTPUT_DIR = "outputs"
os.makedirs(OUTPUT_DIR, exist_ok=True)


def _heading(doc, title):
    p = doc.add_heading(title, level=2)
    p.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(4)


def _bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.runs[0].font.size = Pt(10.5)
    p.paragraph_format.space_after = Pt(2)


def generate_cv_docx(state: dict) -> str:
    """Generates .docx version of the tailored CV. Returns file path."""
    doc      = Document()
    facts    = state["facts_json"]
    personal = facts.get("personal", {})

    # Header
    name_p = doc.add_heading(personal.get("name", ""), level=1)
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    contact = " · ".join(filter(None, [
        personal.get("email", ""),
        personal.get("location", ""),
        f"linkedin.com/{personal.get('linkedin', '')}",
        f"github.com/{personal.get('github', '')}",
    ]))
    cp = doc.add_paragraph(contact)
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.runs[0].font.size = Pt(9.5)

    # Summary
    summary = state.get("tailored_summary", "")
    if summary:
        _heading(doc, "Professional Summary")
        p = doc.add_paragraph(summary)
        p.runs[0].font.size = Pt(10.5)

    # Experience
    experience  = facts.get("experience", [])
    exp_bullets = [b for b in state.get("tailored_bullets", []) if b.get("section") == "experience"]
    if experience:
        _heading(doc, "Experience")
        for exp in experience:
            p = doc.add_paragraph()
            r = p.add_run(f"{exp.get('title', '')}  —  {exp.get('company', '')}")
            r.bold = True
            r.font.size = Pt(10.5)
            dr = p.add_run(f"    {exp.get('dates', '')}")
            dr.font.size = Pt(10)
            dr.font.color.rgb = RGBColor(80, 80, 80)
        for b in exp_bullets:
            _bullet(doc, b["text"])

    # Projects
    projects     = facts.get("projects", [])
    proj_bullets = [b for b in state.get("tailored_bullets", []) if b.get("section") == "project"]
    if projects:
        _heading(doc, "Projects")
        for proj in projects:
            p = doc.add_paragraph()
            r = p.add_run(proj.get("name", ""))
            r.bold = True
            r.font.size = Pt(10.5)
            sr = p.add_run(f"    {', '.join(proj.get('tech_stack', []))}")
            sr.font.size = Pt(10)
            sr.font.color.rgb = RGBColor(80, 80, 80)
        for b in proj_bullets:
            _bullet(doc, b["text"])

    # Skills
    skills = facts.get("skills", {})
    if skills:
        _heading(doc, "Skills")
        for category, items in skills.items():
            if items:
                p = doc.add_paragraph()
                r = p.add_run(f"{category.capitalize()}: ")
                r.bold = True
                r.font.size = Pt(10.5)
                p.add_run(", ".join(items)).font.size = Pt(10.5)
                p.paragraph_format.space_after = Pt(2)

    # Education
    education = facts.get("education", [])
    if education:
        _heading(doc, "Education")
        for edu in education:
            p = doc.add_paragraph()
            r = p.add_run(f"{edu.get('degree', '')}  —  {edu.get('institution', '')}")
            r.bold = True
            r.font.size = Pt(10.5)
            p.add_run(f"    {edu.get('graduation_year', '')}").font.size = Pt(10)
            gp = doc.add_paragraph(f"GPA: {edu.get('gpa', '')}")
            gp.runs[0].font.size = Pt(10)
            gp.runs[0].font.color.rgb = RGBColor(80, 80, 80)
            gp.paragraph_format.space_after = Pt(2)

    output_path = os.path.join(OUTPUT_DIR, "tailored_cv.docx")
    doc.save(output_path)
    logger.info(f"✅ CV DOCX saved → {output_path}")
    return output_path